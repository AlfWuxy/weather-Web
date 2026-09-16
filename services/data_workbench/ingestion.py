"""机构 Excel 安全解析、无患者内容预览和原子追加修订。"""
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime, timedelta

from flask import current_app
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.exc import IntegrityError

from core.extensions import db
from core.pilot_models import PilotImportBatch, PilotEncounter, PilotCoverage, PilotInstitution
from core.time_utils import utcnow, today_local
from .storage import store_bytes, read_bytes, private_digest

ALIASES = {
    'encounter_date': ('挂号日期', '就诊日期', '挂号时间', '就诊时间', '日期', 'date'),
    'age': ('年龄', '患者年龄', 'age'),
    'patient_key': ('病案号', '患者编号', 'patient_key'),
    'source_id': ('就诊流水号', '挂号流水号', '就诊编号', '挂号编号', '就诊ID', 'encounter_id'),
    'residence_region_code': ('居住地行政区代码', '患者居住地区划代码', 'residence_region_code'),
    'diagnosis': ('诊断', '门诊诊断', '诊断名称', '疾病诊断', 'diagnosis'),
}


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for pattern in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%Y年%m月%d日', '%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S'):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                continue
    raise ValueError('日期无效，请使用真实日期')


def _timestamp_key(value):
    # 存在真实时刻时保留到原有精度，仅将格式规范化，不扩展精度。
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, str) and (':' in value):
        try:
            return datetime.fromisoformat(value.strip().replace('/', '-')).isoformat()
        except ValueError:
            return value.strip()
    return parse_date(value).isoformat()


def _age(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('年龄无效')
    if isinstance(value, (float, int)):
        if not math.isfinite(value) or value != int(value):
            raise ValueError('年龄须为完整周岁')
        result = int(value)
    elif isinstance(value, str) and re.fullmatch(r'\s*\d{1,3}\s*岁?\s*', value):
        result = int(re.search(r'\d+', value).group())
    elif isinstance(value, str) and re.fullmatch(r'\s*\d{1,3}岁(?:\d{1,2}月)?(?:\d{1,2}天)?\s*', value):
        match = re.fullmatch(r'\s*(\d{1,3})岁(?:(\d{1,2})月)?(?:(\d{1,2})天)?\s*', value)
        result = int(match.group(1))
        if int(match.group(2) or 0) > 11 or int(match.group(3) or 0) > 31:
            raise ValueError('年龄月数或天数无效')
    elif isinstance(value, str) and re.fullmatch(r'\s*\d{1,3}\s*[月天日]\s*', value):
        number = int(re.search(r'\d+', value).group())
        result = number // (12 if '月' in value else 365)
    else:
        raise ValueError('年龄无效')
    if not 0 <= result <= 120:
        raise ValueError('年龄超出 0 至 120 岁范围')
    return result


def _coverage(coverage):
    start, end = parse_date(coverage.get('start') or coverage.get('coverage_start')), parse_date(coverage.get('end') or coverage.get('coverage_end'))
    mode = coverage.get('mode') or coverage.get('coverage_mode')
    if mode not in {'complete', 'partial'} or end < start or (end - start).days > 3660 or end > today_local():
        raise ValueError('报送范围或完整性无效')
    closed = sorted({parse_date(v) for v in coverage.get('closed_dates', [])})
    if any(d < start or d > end for d in closed):
        raise ValueError('停诊日必须位于报送范围内')
    return start, end, mode, [d.isoformat() for d in closed]


def _safe_workbook(data):
    max_bytes = current_app.config.get('PILOT_MAX_UPLOAD_BYTES', 10 * 1024 * 1024)
    if not data or len(data) > max_bytes:
        raise ValueError('Excel 文件为空或超过上传上限')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(i.file_size for i in entries) > 100 * 1024 * 1024:
                raise ValueError('Excel 解压内容超过上限')
            for entry in entries:
                name = entry.filename.lower()
                if entry.flag_bits & 1 or '..' in name.split('/') or name.startswith('/') or 'vbaproject' in name or 'externallinks/' in name or 'embeddings/' in name:
                    raise ValueError('Excel 包含不支持的嵌入或外部内容')
                if entry.file_size > 20 * 1024 * 1024 or (entry.file_size > 1024 * 1024 and entry.file_size / max(entry.compress_size, 1) > 200):
                    raise ValueError('Excel 压缩比例或单项大小异常')
                if name.endswith('.xml'):
                    xml = archive.read(entry)
                    if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                        raise ValueError('Excel XML 包含不支持的实体定义')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) != 1:
            book.close()
            raise ValueError('请上传只含一张数据工作表的 Excel')
        sheet = book.worksheets[0]
        # 不信任文件声明的维度，实际迭代仍会再次限额。
        if (sheet.max_row or 0) > current_app.config.get('PILOT_MAX_ROWS', 100000) + 1 or (sheet.max_column or 0) > 100:
            book.close()
            raise ValueError('Excel 行列数超过上限')
        return book
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError('文件不是有效的 .xlsx 工作簿') from exc


def _lock_institution(institution_id):
    # 统一先锁机构再锁批次；SQLite 与 PostgreSQL 都能防止同机构并发半批写入。
    db.session.execute(update(PilotInstitution).where(PilotInstitution.id == institution_id).values(enabled=PilotInstitution.enabled))


def _metadata_matches(batch, start, end, mode, closed, mapping):
    return (batch.coverage_start == start and batch.coverage_end == end and batch.coverage_mode == mode
            and sorted(batch.closed_dates or []) == closed
            and (mapping is None or all(batch.mapping.get(k) == v for k, v in mapping.items())))


def create_batch(inst, file_bytes, filename, user_id, coverage, mapping=None):
    if not inst.enabled or not inst.raw_storage_approved or (inst.retention_days or 0) <= 0:
        raise ValueError('机构尚未确认原表存储与保存期限')
    if not str(filename).lower().endswith('.xlsx'):
        raise ValueError('只接受 .xlsx 文件')
    start, end, mode, closed = _coverage(coverage)
    revision_of = coverage.get('revision_of') or None
    mapping = mapping or None
    sha = hashlib.sha256(file_bytes).hexdigest()
    key = hashlib.sha256(f'{sha}:{revision_of or "original"}'.encode()).hexdigest()
    book = _safe_workbook(file_bytes)
    book.close()
    try:
        _lock_institution(inst.id)
        if revision_of:
            old = PilotImportBatch.query.filter_by(id=revision_of, institution_id=inst.id, status='confirmed').first()
            existing = PilotImportBatch.query.filter_by(institution_id=inst.id, idempotency_key=key).first()
            newer = PilotImportBatch.query.filter_by(revision_of=revision_of, institution_id=inst.id, status='confirmed').first()
            if newer and (not existing or newer.id != existing.id):
                raise ValueError('原批次已有更新版本，请选择最新修订批次')
            if not old or old.coverage_start != start or old.coverage_end != end:
                raise ValueError('修订必须选择同机构已确认批次，且覆盖日期保持一致')
        else:
            # 优先找任意已确认版本，未确认的旧尝试不能绕过同文件幂等保护。
            existing = PilotImportBatch.query.filter_by(institution_id=inst.id, sha256=sha, status='confirmed').order_by(PilotImportBatch.confirmed_at.desc()).first()
            if not existing:
                existing = PilotImportBatch.query.filter_by(institution_id=inst.id, idempotency_key=key).first()
        if existing:
            db.session.refresh(existing)
            if existing.status == 'confirmed':
                if not _metadata_matches(existing, start, end, mode, closed, mapping):
                    raise ValueError('相同文件已确认但报送设置不同，请选择原批次进行修订')
                db.session.commit()
                return existing
            if existing.status == 'confirming':
                raise ValueError('批次正在确认，请等待完成后选择修订')
            if not _metadata_matches(existing, start, end, mode, closed, mapping):
                existing.coverage_start, existing.coverage_end = start, end
                existing.coverage_mode, existing.closed_dates = mode, closed
                if mapping is not None:
                    existing.mapping = mapping
                existing.report, existing.status = {}, 'queued'
            db.session.commit()
            return existing
        # 每个修订独立保存原件，避免清理旧批次原件时破坏较新的版本。
        path = store_bytes(file_bytes)
        batch = PilotImportBatch(institution_id=inst.id, user_id=user_id, sha256=sha, idempotency_key=key,
                                 filename='就诊数据.xlsx', storage_path=path, coverage_start=start,
                                 coverage_end=end, coverage_mode=mode, closed_dates=closed,
                                 mapping=mapping if mapping is not None else (inst.field_mapping or {}), revision_of=revision_of)
        db.session.add(batch)
        db.session.commit()
        return batch
    except Exception:
        db.session.rollback()
        raise


def _category(value):
    text = str(value or '').strip()
    if not text:
        return 'unknown'
    if any(s in text for s in ('呼吸', '肺炎', '咳嗽', '感冒', '哮喘', '气管')):
        return 'respiratory'
    if any(s in text for s in ('血压', '心脏', '冠心', '心衰', '脑梗', '心律')):
        return 'cardiovascular'
    return 'other'


def _parse(batch):
    book = _safe_workbook(read_bytes(batch.storage_path))
    try:
        iterator = book.worksheets[0].iter_rows()
        cells = next(iterator, ())
        headers = [str(c.value or '').strip() for c in cells]
        if not headers or len(headers) != len(set(headers)) or any(not h for h in headers):
            raise ValueError('第一行须为不重复且非空的字段名')
        if any(len(h) > 80 or c.data_type == 'f' for h, c in zip(headers, cells)):
            raise ValueError('字段名过长或包含公式')
        mapping = dict(batch.mapping or {})
        if set(mapping) - set(ALIASES):
            raise ValueError('包含不支持的字段映射')
        for key, aliases in ALIASES.items():
            if key not in mapping:
                match = [h for h in headers if h in aliases]
                if len(match) == 1:
                    mapping[key] = match[0]
        if any(mapping.get(k) not in headers for k in ('encounter_date', 'age')):
            return [], {'headers': headers, 'mapping': mapping, 'needs_mapping': True, 'counts': {'rows': 0, 'valid': 0, 'invalid': 0, 'duplicates': 0, 'overlap': 0, 'missing_diagnosis': 0}, 'issues': [], 'duplicate_groups': [], 'revision_rows': [], 'revision_of': batch.revision_of}
        if any(v and v not in headers for v in mapping.values()) or len([v for v in mapping.values() if v]) != len(set(v for v in mapping.values() if v)):
            raise ValueError('字段映射不存在或重复使用同一列')
        indices = {k: headers.index(v) for k, v in mapping.items() if v}
        parsed, issues, total, missing = [], [], 0, 0
        for row_no, row in enumerate(iterator, 2):
            if row_no > current_app.config.get('PILOT_MAX_ROWS', 100000) + 1 or len(row) > 100:
                raise ValueError('Excel 行列数超过上限')
            if all(c.value is None for c in row):
                continue
            total += 1
            try:
                if any(c.data_type == 'f' for c in row):
                    raise ValueError('包含公式，请导出固定值')
                values = {k: row[index].value for k, index in indices.items()}
                day = parse_date(values['encounter_date'])
                age = _age(values['age'])
                if not batch.coverage_start <= day <= batch.coverage_end or day.isoformat() in batch.closed_dates:
                    raise ValueError('就诊日期超出报送范围或位于停诊日')
                diagnosis = str(values.get('diagnosis') or '').strip()
                if not diagnosis:
                    missing += 1
                source_id = str(values.get('source_id') or '').strip()
                residence = str(values.get('residence_region_code') or '').strip()
                if residence and not re.fullmatch(r'\d{6}|\d{9}|\d{12}', residence):
                    raise ValueError('居住地仅接受 6、9 或 12 位行政区代码，不接收详细地址')
                patient_key = str(values.get('patient_key') or '').strip()
                suspected = private_digest(batch.institution_id, json.dumps(['suspected', patient_key, _timestamp_key(values['encounter_date'])], ensure_ascii=False)) if patient_key else None
                fp = private_digest(batch.institution_id, json.dumps([_timestamp_key(values['encounter_date']), age, diagnosis, residence, str(values.get('patient_key') or '').strip()], ensure_ascii=False))
                parsed.append({'row': row_no, 'date': day, 'age': age, 'source_key': private_digest(batch.institution_id, 'id:' + source_id) if source_id else None,
                               'fingerprint': fp, 'suspected_key': suspected, 'diagnosis_category': _category(diagnosis), 'residence_region_code': residence or None})
            except (ValueError, IndexError) as exc:
                issues.append({'row': row_no, 'message': str(exc) if isinstance(exc, ValueError) else '行数据列数不足'})
        counts = {'rows': total, 'valid': len(parsed), 'invalid': len(issues), 'duplicates': 0, 'overlap': 0, 'missing_diagnosis': missing}
        report = {'headers': headers, 'mapping': mapping, 'needs_mapping': False, 'counts': counts,
                  'issues': issues[:200], 'issues_truncated': len(issues) > 200, 'duplicate_groups': [], 'revision_rows': [], 'revision_of': batch.revision_of,
                  'date_start': min((r['date'] for r in parsed), default=batch.coverage_start).isoformat(),
                  'date_end': max((r['date'] for r in parsed), default=batch.coverage_end).isoformat()}
        return parsed, report
    finally:
        book.close()


def _analyze(batch, rows, report):
    existing = PilotEncounter.query.filter_by(institution_id=batch.institution_id, active=True).all()
    if batch.revision_of:
        existing = [r for r in existing if r.batch_id != batch.revision_of]
    source = {r.source_key: r for r in existing if r.source_key}
    fingerprints, suspected = {}, {}
    for record in existing:
        fingerprints.setdefault(record.fingerprint, set()).add(record.id)
        if record.suspected_key:
            suspected.setdefault(record.suspected_key, set()).add(record.id)
    seen_ids, groups = {}, {}
    for row in rows:
        key = row['source_key']
        if key:
            if key in seen_ids and seen_ids[key]['fingerprint'] != row['fingerprint']:
                report['issues'].append({'row': row['row'], 'message': '同一就诊编号在文件内内容冲突，请修订原表'})
                report['counts']['invalid'] += 1
            elif key in seen_ids:
                row['automatic_skip'] = True
                report['counts']['duplicates'] += 1
            seen_ids[key] = row
            if key in source:
                report['counts']['overlap'] += 1
                if source[key].fingerprint == row['fingerprint']:
                    row['automatic_skip'] = True
                else:
                    row['supersedes_id'] = source[key].id
                    report['revision_rows'].append(row['row'])
        else:
            groups.setdefault(row.get('suspected_key') or row['fingerprint'], []).append(row)
    for group_key, group in groups.items():
        matches = set()
        for row in group:
            matches.update(fingerprints.get(row['fingerprint'], set()))
            matches.update(suspected.get(row.get('suspected_key'), set()))
        count = len(matches)
        if len(group) > 1 or count:
            report['duplicate_groups'].append({'id': group_key, 'rows': [r['row'] for r in group], 'existing_count': count})
            report['counts']['duplicates'] += len(group) - (0 if count else 1)
            report['counts']['overlap'] += len(group) if count else 0
    return rows, report


def process_batch(batch_id):
    batch = db.session.get(PilotImportBatch, batch_id)
    if not batch:
        raise ValueError('导入批次不存在')
    try:
        _lock_institution(batch.institution_id)
        db.session.refresh(batch)
        if batch.status == 'confirmed':
            db.session.commit()
            return batch.report
        rows, report = _parse(batch)
        _analyze(batch, rows, report)
        batch.report, batch.mapping = report, report['mapping']
        batch.status = 'needs_mapping' if report['needs_mapping'] else ('invalid' if report['counts']['invalid'] else 'ready')
        db.session.commit()
        return report
    except Exception:
        db.session.rollback()
        raise


def confirm_batch(batch_id, user_id, decisions=None):
    decisions = decisions or {}
    batch = db.session.get(PilotImportBatch, batch_id)
    if not batch:
        raise ValueError('导入批次不存在')
    try:
        _lock_institution(batch.institution_id)
        db.session.refresh(batch)
        if batch.status == 'confirmed':
            db.session.commit()
            return batch.report
        # 在机构锁内条件更新批次；失败全部回滚，不允许半批入库。
        changed = db.session.execute(update(PilotImportBatch).where(PilotImportBatch.id == batch_id, PilotImportBatch.status == 'ready').values(status='confirming')).rowcount
        if changed != 1:
            raise ValueError('批次尚未通过检查或正在确认')
        inst = db.session.execute(db.select(PilotInstitution).where(PilotInstitution.id == batch.institution_id).with_for_update()).scalar_one()
        rows, report = _parse(batch)
        rows, report = _analyze(batch, rows, report)
        if report['needs_mapping'] or report['counts']['invalid']:
            raise ValueError('请先修正字段或异常行')
        if report['counts']['missing_diagnosis'] and decisions.get('accept_missing_diagnosis') is not True:
            raise ValueError('存在缺失诊断，须确认按未知分类保留')
        revisions = {int(v) for v in decisions.get('revision_rows', [])}
        if set(report['revision_rows']) - revisions:
            raise ValueError('同一就诊编号存在修订，须逐行确认替代')
        if batch.revision_of and PilotImportBatch.query.filter(PilotImportBatch.revision_of == batch.revision_of, PilotImportBatch.status == 'confirmed', PilotImportBatch.id != batch.id).first():
            raise ValueError('原批次已有更新版本，请选择最新修订批次')
        if batch.revision_of and decisions.get('confirm_revision') is not True:
            raise ValueError('须确认此批次替代原批次记录')
        skip = {r['row'] for r in rows if r.get('automatic_skip')}
        for group in report['duplicate_groups']:
            choice = decisions.get('duplicates', {}).get(group['id'])
            if choice not in {'keep_first', 'keep_all', 'skip'}:
                raise ValueError('请逐组确认疑似重复记录')
            if choice == 'skip' or (choice == 'keep_first' and group['existing_count']):
                skip.update(group['rows'])
            elif choice == 'keep_first':
                skip.update(group['rows'][1:])
        if batch.revision_of:
            db.session.execute(update(PilotEncounter).where(PilotEncounter.batch_id == batch.revision_of, PilotEncounter.active.is_(True)).values(active=False))
        for row in rows:
            if row['row'] in skip:
                continue
            supersedes = row.get('supersedes_id')
            if supersedes:
                db.session.get(PilotEncounter, supersedes).active = False
                db.session.flush()
            db.session.add(PilotEncounter(institution_id=batch.institution_id, batch_id=batch.id,
                encounter_date=row['date'], age=row['age'], source_key=row['source_key'], fingerprint=row['fingerprint'], suspected_key=row['suspected_key'],
                diagnosis_category=row['diagnosis_category'], residence_region_code=row['residence_region_code'], supersedes_id=supersedes))
        db.session.flush()
        day = batch.coverage_start
        while day <= batch.coverage_end:
            status = 'closed' if day.isoformat() in batch.closed_dates else batch.coverage_mode
            if status == 'closed' and PilotEncounter.query.filter_by(institution_id=batch.institution_id, encounter_date=day, active=True).first():
                raise ValueError('停诊日仍有有效就诊记录，请核对修订范围')
            previous = PilotCoverage.query.filter_by(institution_id=batch.institution_id, date=day).first()
            if previous:
                # 追加部分报送不会撤销原先确认的完整性；整批修订允许重新声明。
                if batch.revision_of or status == 'closed' or previous.status != 'complete' or status == 'complete':
                    previous.status, previous.batch_id = status, batch.id
            else:
                db.session.add(PilotCoverage(institution_id=batch.institution_id, date=day, batch_id=batch.id, status=status))
            day += timedelta(days=1)
        report['imported_rows'] = len(rows) - len(skip)
        report['skipped_rows'] = len(skip)
        batch.report, batch.status, batch.confirmed_at = report, 'confirmed', utcnow()
        inst.field_mapping = report['mapping']
        db.session.commit()
        return report
    except Exception:
        db.session.rollback()
        raise
